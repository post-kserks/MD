"""Собирает report.docx из report.md (с титульным листом).

Нужен python-docx:  pip install python-docx
Запуск:             python3 build_docx.py
Поддерживается только то подмножество Markdown, которое используется в report.md.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
FONT = "Times New Roman"


def set_cell_borders(table):
    tbl = table._tbl
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl.tblPr.append(borders)


def add_runs(par, text, size=None, bold=False):
    """Разбирает **жирный**, *курсив* и `код`."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*"):
            run = par.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
        else:
            run = par.add_run(part.replace("\\", ""))
        if bold:
            run.bold = True
        if size:
            run.font.size = Pt(size)


def para(doc, text="", align=None, size=None, bold=False, indent=True, space_after=6):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    add_runs(p, text, size, bold)
    return p


def title_page(doc):
    c = WD_ALIGN_PARAGRAPH.CENTER
    for line in (
        "Министерство науки и высшего образования Российской Федерации",
        "Федеральное государственное автономное образовательное учреждение высшего образования",
        "«Московский государственный технический университет имени Н.Э. Баумана",
        "(национальный исследовательский университет)»",
        "(МГТУ им. Н.Э. Баумана)",
    ):
        para(doc, line, c, 12, indent=False, space_after=0)
    para(doc, "", indent=False)
    para(doc, "Факультет «Информатика и системы управления»", c, 13, indent=False, space_after=0)
    para(doc, "Кафедра «Системы обработки информации и управления»", c, 13, indent=False)
    for _ in range(5):
        para(doc, "", indent=False, space_after=0)
    para(doc, "Отчёт по лабораторной работе № 1 (часть 1)", c, 16, True, indent=False)
    para(doc, "по дисциплине «Модели данных»", c, 14, indent=False)
    para(doc, "«Исследование предметной области с использованием ИИ и построение инфологической модели»", c, 14, indent=False)
    para(doc, "Тема: «АС учёта компьютеров, используемых фирмой»", c, 14, True, indent=False)
    for _ in range(5):
        para(doc, "", indent=False, space_after=0)
    r = WD_ALIGN_PARAGRAPH.RIGHT
    para(doc, "Выполнил: студент группы ________", r, 14, indent=False, space_after=0)
    para(doc, "________________ (ФИО)", r, 14, indent=False)
    para(doc, "Проверил: ________________", r, 14, indent=False)
    for _ in range(4):
        para(doc, "", indent=False, space_after=0)
    para(doc, "Москва, 2026", c, 14, indent=False)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(t)
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], h, 11, bold=True)
    for row in body:
        cells = t.add_row().cells
        for i, v in enumerate(row[: len(header)]):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], v, 11)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main():
    lines = (HERE / "report.md").read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.line_spacing = 1.15
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        hs = doc.styles[name]
        hs.font.name = FONT
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        hs.font.color.rgb = None

    title_page(doc)
    content_w = sec.page_width - sec.left_margin - sec.right_margin

    i = 0
    # пропускаем заголовок документа — он на титульном листе
    while i < len(lines) and not lines[i].startswith("## 1."):
        i += 1
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|[\s:\-|]+\|$", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"^(#{2,4})\s+(.*)", s)
        if m:
            level = len(m.group(1)) - 1
            h = doc.add_heading(level=level)
            add_runs(h, m.group(2).replace("**", ""))
            for r in h.runs:
                r.font.size = Pt({1: 16, 2: 15, 3: 14}[level])
            i += 1
            continue
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", s)
        if m:
            doc.add_picture(str(HERE / m.group(1)), width=content_w)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue
        if s.startswith(">"):
            p = para(doc, s.lstrip("> ").strip(), indent=False)
            p.paragraph_format.left_indent = Cm(1)
            for r in p.runs:
                r.italic = True
            i += 1
            continue
        if s.startswith("*Рисунок"):
            para(doc, s.strip("*"), WD_ALIGN_PARAGRAPH.CENTER, 12, indent=False)
            i += 1
            continue
        m = re.match(r"^[-•]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)", s)
        if m:
            p = para(doc, f"{m.group(1)}. {m.group(2)}", indent=False, space_after=3)
            p.paragraph_format.left_indent = Cm(0.75)
            i += 1
            continue
        para(doc, s)
        i += 1

    out = HERE / "report.docx"
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
